import io
import csv
import json
from datetime import datetime
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database import Registration

def generate_excel_spreadsheet(registrations: List[Registration]) -> bytes:
    """
    Builds a styled Excel spreadsheet containing all registered teams
    with full leader and member breakdowns.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ANTROTHON'27 Registrations"

    # Define Styles
    header_fill = PatternFill(start_color="1F1113", end_color="1F1113", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FF3B30")
    
    sub_fill = PatternFill(start_color="2D181A", end_color="2D181A", fill_type="solid")
    sub_font = Font(name="Calibri", size=10, bold=True, color="F2F2F7")
    
    paid_fill = PatternFill(start_color="1C3829", end_color="1C3829", fill_type="solid")
    paid_font = Font(name="Calibri", size=10, bold=True, color="34C759")
    
    pending_fill = PatternFill(start_color="3D2A12", end_color="3D2A12", fill_type="solid")
    pending_font = Font(name="Calibri", size=10, bold=True, color="FF9F0A")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # Title row
    ws.merge_cells('A1:AC1')
    title_cell = ws['A1']
    title_cell.value = "ANTROTHON'27 — INTER-COLLEGE HACKATHON REGISTRATIONS"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0B0C10", end_color="0B0C10", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    ws.merge_cells('A2:AC2')
    sub_cell = ws['A2']
    sub_cell.value = f"Patrician Institute Innovation Cell (PIIC) | Generated: {datetime.now().strftime('%d-%b-%Y %I:%M %p')} | Total Teams: {len(registrations)}"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="E5E5EA")
    sub_cell.fill = PatternFill(start_color="161B22", end_color="161B22", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    headers = [
        "S.No", "Team Code", "Team Name", "Track", "Status", 
        "Payment Status", "Amount (INR)", "Payment ID", "Order ID",
        "Leader Name", "Leader Email", "Leader Phone", "Leader College", 
        "Leader Department", "Leader Year", "Leader Gender", "Leader Student ID",
        "Team Size",
        # Member 2
        "M2 Name", "M2 Email", "M2 Phone", "M2 College", "M2 Dept",
        # Member 3
        "M3 Name", "M3 Email", "M3 Phone", "M3 College", "M3 Dept",
        # Member 4
        "M4 Name", "M4 Email", "M4 Phone", "M4 College", "M4 Dept",
        # Member 5 (Optional)
        "M5 Name", "M5 Email", "M5 Phone", "M5 College", "M5 Dept",
        # Member 6 (Optional)
        "M6 Name", "M6 Email", "M6 Phone", "M6 College", "M6 Dept",
        "Idea Title", "Abstract / Summary", "Idea Description", "PPT Presentation File", "Registered At (UTC)"
    ]

    ws.append([])  # Row 3 is blank spacer
    ws.append(headers)  # Row 4 is Headers
    ws.row_dimensions[4].height = 28

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Populate Data
    for idx, reg in enumerate(registrations, 1):
        members = reg.get_members()
        m2 = members[0] if len(members) > 0 else {}
        m3 = members[1] if len(members) > 1 else {}
        m4 = members[2] if len(members) > 2 else {}
        m5 = members[3] if len(members) > 3 else {}
        m6 = members[4] if len(members) > 4 else {}

        row_data = [
            idx,
            reg.team_code,
            reg.team_name,
            reg.track,
            reg.status.replace("_", " ").title(),
            reg.payment_status.upper(),
            reg.amount_paid,
            reg.razorpay_payment_id or "-",
            reg.razorpay_order_id or "-",
            reg.leader_name,
            reg.leader_email,
            reg.leader_phone,
            reg.college_name or reg.leader_college,
            reg.leader_department,
            reg.leader_year,
            reg.leader_gender or "-",
            reg.leader_id_number or "-",
            reg.member_count,
            # M2
            m2.get("name", "-"), m2.get("email", "-"), m2.get("phone", "-"), m2.get("college", "-"), m2.get("department", "-"),
            # M3
            m3.get("name", "-"), m3.get("email", "-"), m3.get("phone", "-"), m3.get("college", "-"), m3.get("department", "-"),
            # M4
            m4.get("name", "-"), m4.get("email", "-"), m4.get("phone", "-"), m4.get("college", "-"), m4.get("department", "-"),
            # M5
            m5.get("name", "-"), m5.get("email", "-"), m5.get("phone", "-"), m5.get("college", "-"), m5.get("department", "-"),
            # M6
            m6.get("name", "-"), m6.get("email", "-"), m6.get("phone", "-"), m6.get("college", "-"), m6.get("department", "-"),
            reg.idea_title or reg.project_title or "-",
            reg.idea_summary or reg.project_synopsis or "-",
            reg.idea_description or "-",
            reg.idea_ppt_filename or "-",
            reg.created_at.strftime("%Y-%m-%d %H:%M:%S") if reg.created_at else "-"
        ]

        ws.append(row_data)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 22

        # Format row cells
        for col_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.border = thin_border
            c.alignment = Alignment(vertical="center")

            # Status highlights
            if col_idx == 6: # Payment status
                if str(c.value).upper() == "PAID":
                    c.fill = paid_fill
                    c.font = paid_font
                else:
                    c.fill = pending_fill
                    c.font = pending_font
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (1, 2, 5, 7, 18):
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2):
                continue
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_csv_spreadsheet(registrations: List[Registration]) -> str:
    """
    Builds a CSV string for export.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Team Code", "Team Name", "Track", "Status", "Payment Status", "Amount INR",
        "Payment ID", "Leader Name", "Leader Email", "Leader Phone", "College",
        "Leader Department", "Leader Year", "Member Count", "Members Details",
        "Idea Title", "Abstract / Summary", "Idea Description", "PPT Presentation File", "Registered At"
    ])

    for reg in registrations:
        members_summary = "; ".join([
            f"{m.get('name')} ({m.get('department', m.get('dept', ''))}, {m.get('phone', '')}, {m.get('email', '')})"
            for m in reg.get_members()
        ])
        writer.writerow([
            reg.team_code,
            reg.team_name,
            reg.track,
            reg.status,
            reg.payment_status,
            reg.amount_paid,
            reg.razorpay_payment_id or "",
            reg.leader_name,
            reg.leader_email,
            reg.leader_phone,
            reg.college_name or reg.leader_college,
            reg.leader_department,
            reg.leader_year,
            reg.member_count,
            members_summary,
            reg.idea_title or reg.project_title or "",
            reg.idea_summary or reg.project_synopsis or "",
            reg.idea_description or "",
            reg.idea_ppt_filename or "",
            reg.created_at.strftime("%Y-%m-%d %H:%M:%S") if reg.created_at else ""
        ])

    return output.getvalue()

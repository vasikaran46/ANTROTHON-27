import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_template():
    downloads_dir = Path(__file__).resolve().parent.parent / "app" / "static" / "downloads"
    downloads_dir.mkdir(parents=True, exist_ok=True)
    output_path = downloads_dir / "ANTROTHON27_Idea_Presentation_Template.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    COLOR_DARK = RGBColor(10, 14, 20)
    COLOR_RED = RGBColor(255, 30, 46)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_STEEL = RGBColor(148, 163, 184)
    COLOR_CARD = RGBColor(18, 24, 36)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_DARK

    def add_header(slide, title_text, category_text="// ANTROTHON'27 IDEA DOSSIER"):
        # Category / Eyebrow
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_RED

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_WHITE

    blank_layout = prs.slide_layouts[6]

    # --- SLIDE 1: Title Slide ---
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    tbox1 = slide1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(5.8))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "ANTROTHON'27 // IDEA SUBMISSION TEMPLATE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "[ ENTER YOUR IDEA TITLE HERE (MAX 100 CHARS) ]"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_after = Pt(16)

    p = tf1.add_paragraph()
    p.text = "Track: [ Select: Track 1 (Human Safety) / Track 2 (Health & Education) / Track 3 (Sustainability) ]"
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_STEEL
    p.space_after = Pt(24)

    p = tf1.add_paragraph()
    p.text = "Team Name: [ Enter Team Name ]   |   College: [ Enter College Name ]"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_after = Pt(12)

    p = tf1.add_paragraph()
    p.text = "• Team Leader: [ Name | Mobile | Email ]\n• Member 2: [ Name | Mobile | Email ]\n• Member 3: [ Name | Mobile | Email ]\n• Member 4: [ Name | Mobile | Email ]\n• Member 5 (Optional): [ Name | Mobile | Email ]\n• Member 6 (Optional): [ Name | Mobile | Email ]"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_STEEL

    # --- SLIDE 2: Track Focus & Proposed Solution ---
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "01 // Track Focus & Proposed Solution")

    tbox2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf2 = tbox2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "A. Track Focus & Need Addressed:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf2.add_paragraph()
    p.text = "• Specify your chosen Hackathon Track (Track 1, 2, or 3).\n• Describe the societal challenge or opportunity targeted within this track."
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_WHITE
    p.space_after = Pt(18)

    p = tf2.add_paragraph()
    p.text = "B. Proposed Innovative Solution:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf2.add_paragraph()
    p.text = "• Detail your concept and how it directly addresses the targeted track theme.\n• Explain what makes your approach creative, unique, and human-centric.\n• What are the core deliverables of this project?"
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_WHITE

    # --- SLIDE 3: Technical Architecture & Tech Stack ---
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "02 // Technical Architecture & Implementation")

    tbox3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf3 = tbox3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "A. System Architecture / Flow Diagram:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf3.add_paragraph()
    p.text = "[ Insert your system architecture diagram, block diagram, or user workflow here ]"
    p.font.size = Pt(14)
    p.font.italic = True
    p.font.color.rgb = COLOR_STEEL
    p.space_after = Pt(20)

    p = tf3.add_paragraph()
    p.text = "B. Technologies & Tools Planned:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf3.add_paragraph()
    p.text = "• Frontend / Interface: [ e.g. React, Flutter, Web, Android ]\n• Backend & APIs: [ e.g. Python, FastAPI, Node.js, Go ]\n• AI / ML / Models: [ e.g. PyTorch, OpenCV, LLMs, Edge AI ]\n• Database & Hardware: [ e.g. PostgreSQL, Raspberry Pi, ESP32, Firebase ]"
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_WHITE

    # --- SLIDE 4: Feasibility, Viability & Novelty ---
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "03 // Feasibility, Viability & Novelty")

    tbox4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf4 = tbox4.text_frame
    tf4.word_wrap = True
    p = tf4.paragraphs[0]
    p.text = "A. Novelty / Uniqueness (USP):"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf4.add_paragraph()
    p.text = "• How does your solution compare to existing products, open-source projects, or government apps?\n• What is the critical innovation that sets your team apart?"
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_WHITE
    p.space_after = Pt(18)

    p = tf4.add_paragraph()
    p.text = "B. Feasibility & 24-Hour Hackathon Scope:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf4.add_paragraph()
    p.text = "• What prototype will your team build during the 24-hour on-campus hackathon on Jan 7–8, 2027?\n• What are potential technical risks, and how will you mitigate them?"
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_WHITE

    # --- SLIDE 5: Societal Impact & Commercial Potential ---
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "04 // Societal Impact & Target Beneficiaries")

    tbox5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf5 = tbox5.text_frame
    tf5.word_wrap = True
    p = tf5.paragraphs[0]
    p.text = "A. Social Impact & Benefits:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf5.add_paragraph()
    p.text = "• Who are the direct beneficiaries (e.g. students, women, farmers, marginalized communities, hospitals)?\n• Quantifiable impact metrics (e.g. reduction in response time, accessibility increase)."
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_WHITE
    p.space_after = Pt(18)

    p = tf5.add_paragraph()
    p.text = "B. Future Roadmap & Sustainability:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(6)

    p = tf5.add_paragraph()
    p.text = "• Can this project scale into a startup or community deployment via PIIC incubator?\n• Summary and closing remarks."
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_WHITE

    prs.save(str(output_path))
    print(f"[+] PPT Template generated successfully at: {output_path}")

if __name__ == "__main__":
    create_template()
